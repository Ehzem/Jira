import io
import os
from pathlib import Path
from datetime import date, timedelta

import pandas as pd
import plotly.express as px
import requests
import streamlit as st
from openpyxl.styles import Font, PatternFill, Border, Side
from requests.auth import HTTPBasicAuth


st.set_page_config(
    page_title="Jira Student Capture Dashboard",
    page_icon="📊",
    layout="wide",
)


# ---------------------------------------------------------------------------
# ROSTER
# ---------------------------------------------------------------------------

def find_roster_file():
    """
    Find the student repository workbook automatically.

    Priority:
    1. STUDENT_REPO_PATH environment variable
    2. student_repo.xlsx beside this script
    3. Any student_repo*.xlsx beside this script
    """
    configured = os.getenv("STUDENT_REPO_PATH", "").strip()
    if configured:
        configured_path = Path(configured)
        if configured_path.exists():
            return configured_path

    script_dir = Path(__file__).resolve().parent

    preferred = script_dir / "student_repo.xlsx"
    if preferred.exists():
        return preferred

    matches = sorted(script_dir.glob("student_repo*.xlsx"))
    if matches:
        return matches[0]

    return None


@st.cache_data(show_spinner=False)
def load_student_roster(path_string):
    """
    Load student name, email, section and Jira URL from the Excel roster.
    """
    roster = pd.read_excel(path_string)

    # Normalise column names from the supplied workbook.
    roster.columns = [str(c).strip() for c in roster.columns]

    required = ["Name", "Email", "Section", "Jira URL"]
    missing = [c for c in required if c not in roster.columns]
    if missing:
        raise RuntimeError(
            "Student roster is missing required columns: "
            + ", ".join(missing)
        )

    roster = roster[required + (["id"] if "id" in roster.columns else [])].copy()

    for col in ["Name", "Email", "Section", "Jira URL"]:
        roster[col] = roster[col].fillna("").astype(str).str.strip()

    roster = roster[
        (roster["Name"] != "")
        & (roster["Jira URL"] != "")
    ].copy()

    roster["Jira URL"] = roster["Jira URL"].str.rstrip("/")

    # Unique internal key prevents problems if two students have the same name.
    roster["student_key"] = (
        roster["Name"]
        + " | "
        + roster["Email"]
        + " | "
        + roster["Jira URL"]
    )

    return roster.reset_index(drop=True)


# ---------------------------------------------------------------------------
# JIRA
# ---------------------------------------------------------------------------

def jira_request(session, method, url, **kwargs):
    response = session.request(method, url, timeout=30, **kwargs)

    if response.status_code == 401:
        raise RuntimeError("Authentication failed")
    if response.status_code == 403:
        raise RuntimeError("Access denied")
    if response.status_code == 404:
        raise RuntimeError("Jira site or API endpoint not found")

    if not response.ok:
        try:
            details = response.json()
        except ValueError:
            details = response.text
        raise RuntimeError(
            f"Jira returned HTTP {response.status_code}: {details}"
        )

    return response


def build_jql(start_date, end_date, issue_type=None):
    """
    Query issue creation over the selected date range.
    End date is inclusive in the dashboard and exclusive in Jira JQL.
    """
    end_exclusive = end_date + timedelta(days=1)

    clauses = [
        f'created >= "{start_date:%Y-%m-%d}"',
        f'created < "{end_exclusive:%Y-%m-%d}"',
    ]

    if issue_type and issue_type != "All issue types":
        escaped = issue_type.replace('"', '\\"')
        clauses.append(f'issuetype = "{escaped}"')

    return " AND ".join(clauses) + " ORDER BY created ASC"


def fetch_student_issues(
    email,
    api_token,
    student_name,
    student_email,
    section,
    jira_url,
    jql,
):
    """
    Read issue creation data from one student's Jira site.

    The student's roster entry is treated as the unit being graded.
    This means the report attributes issue creation on that Jira site to
    the student mapped to that site in the roster.
    """
    session = requests.Session()
    session.auth = HTTPBasicAuth(email, api_token)

    url = f"{jira_url}/rest/api/3/search/jql"
    issues = []
    next_page_token = None

    while True:
        body = {
            "jql": jql,
            "maxResults": 100,
            "fields": [
                "summary",
                "created",
                "issuetype",
                "project",
                "status",
                "creator",
            ],
        }

        if next_page_token:
            body["nextPageToken"] = next_page_token

        payload = jira_request(
            session,
            "POST",
            url,
            headers={
                "Accept": "application/json",
                "Content-Type": "application/json",
            },
            json=body,
        ).json()

        batch = payload.get("issues", [])
        issues.extend(batch)

        if payload.get("isLast", False):
            break

        next_page_token = payload.get("nextPageToken")
        if not next_page_token:
            break

    rows = []

    for issue in issues:
        fields = issue.get("fields") or {}
        project = fields.get("project") or {}
        issue_type = fields.get("issuetype") or {}
        status = fields.get("status") or {}
        creator = fields.get("creator") or {}

        rows.append(
            {
                "student_name": student_name,
                "student_email": student_email,
                "section": section,
                "jira_url": jira_url,
                "key": issue.get("key"),
                "summary": fields.get("summary"),
                "created": fields.get("created"),
                "project_key": project.get("key"),
                "project_name": project.get("name"),
                "issue_type": issue_type.get("name"),
                "status": status.get("name"),
                "actual_creator": (
                    creator.get("displayName") or "Unknown creator"
                ),
            }
        )

    df = pd.DataFrame(rows)

    if not df.empty:
        df["created"] = (
            pd.to_datetime(df["created"], utc=True)
            .dt.tz_convert(None)
        )

    return df


# ---------------------------------------------------------------------------
# PERIOD / METRIC LOGIC
# ---------------------------------------------------------------------------

def get_bucket_start(series, granularity):
    normalized = series.dt.normalize()

    if granularity == "Daily":
        return normalized

    week_start = (
        normalized
        - pd.to_timedelta(series.dt.weekday, unit="D")
    )

    if granularity == "Weekly":
        return week_start

    anchor = pd.Timestamp("1970-01-05")
    weeks_from_anchor = (
        (week_start - anchor).dt.days // 7
    ).astype(int)
    biweek_index = weeks_from_anchor // 2

    return (
        anchor
        + pd.to_timedelta(biweek_index * 14, unit="D")
    )


def empty_bucket_index(start_date, end_date, granularity):
    all_days = pd.Series(
        pd.date_range(
            pd.Timestamp(start_date),
            pd.Timestamp(end_date),
            freq="D",
        )
    )

    if granularity == "Daily":
        buckets = all_days.dt.normalize()

    elif granularity == "Weekly":
        buckets = (
            all_days.dt.normalize()
            - pd.to_timedelta(
                all_days.dt.weekday,
                unit="D",
            )
        )

    else:
        week_start = (
            all_days.dt.normalize()
            - pd.to_timedelta(
                all_days.dt.weekday,
                unit="D",
            )
        )

        anchor = pd.Timestamp("1970-01-05")
        weeks_from_anchor = (
            (week_start - anchor).dt.days // 7
        ).astype(int)
        biweek_index = weeks_from_anchor // 2

        buckets = (
            anchor
            + pd.to_timedelta(
                biweek_index * 14,
                unit="D",
            )
        )

    return pd.DatetimeIndex(
        sorted(pd.unique(buckets))
    )


def build_student_summary(
    issues_df,
    selected_students_df,
    start_date,
    end_date,
    granularity,
):
    """
    Build one row for every selected student, including students with zero
    created issues.
    """
    total_periods = len(
        empty_bucket_index(
            start_date,
            end_date,
            granularity,
        )
    )

    working = issues_df.copy()

    if not working.empty:
        working["period_start"] = get_bucket_start(
            working["created"],
            granularity,
        )

    rows = []

    for _, student in selected_students_df.iterrows():
        student_name = student["Name"]
        student_email = student["Email"]
        section = student["Section"]
        jira_url = student["Jira URL"]

        if working.empty:
            student_issues = working
        else:
            student_issues = working[
                working["student_name"] == student_name
            ].copy()

        issues_created = len(student_issues)

        if student_issues.empty:
            active_periods = 0
        else:
            active_periods = int(
                student_issues["period_start"].nunique()
            )

        consistency = (
            active_periods / total_periods * 100
            if total_periods
            else 0
        )

        capture_grade = consistency / 100 * 4

        average_per_period = (
            issues_created / total_periods
            if total_periods
            else 0
        )

        rows.append(
            {
                "Student": student_name,
                "Issues created": issues_created,
                "Capture consistency": f"{consistency:.1f}%",
                "Capture grade": round(capture_grade, 2),
                "Average per period": round(
                    average_per_period,
                    1,
                ),
                "Periods with activity": (
                    f"{active_periods} / {total_periods}"
                ),
                "Email": student_email,
                "Section": section,
                "Jira URL": jira_url,
                "_consistency_numeric": consistency,
            }
        )

    return pd.DataFrame(rows)


def build_chart_data(
    issues_df,
    selected_students_df,
    start_date,
    end_date,
    granularity,
):
    bucket_index = empty_bucket_index(
        start_date,
        end_date,
        granularity,
    )

    students = selected_students_df["Name"].tolist()

    if not students:
        return pd.DataFrame(
            columns=[
                "period_start",
                "student_name",
                "issues_created",
                "period",
            ]
        )

    if issues_df.empty:
        grouped = pd.Series(
            dtype="int64",
            name="issues_created",
        )
    else:
        working = issues_df.copy()
        working["period_start"] = get_bucket_start(
            working["created"],
            granularity,
        )

        grouped = (
            working.groupby(
                ["period_start", "student_name"]
            )
            .size()
            .rename("issues_created")
        )

    full_index = pd.MultiIndex.from_product(
        [bucket_index, students],
        names=["period_start", "student_name"],
    )

    chart_df = (
        grouped
        .reindex(full_index, fill_value=0)
        .reset_index()
    )

    if granularity == "Daily":
        chart_df["period"] = (
            chart_df["period_start"]
            .dt.strftime("%d %b %Y")
        )

    elif granularity == "Weekly":
        period_end = (
            chart_df["period_start"]
            + pd.Timedelta(days=6)
        )

        chart_df["period"] = (
            chart_df["period_start"]
            .dt.strftime("%d %b")
            + " – "
            + period_end.dt.strftime("%d %b %Y")
        )

    else:
        period_end = (
            chart_df["period_start"]
            + pd.Timedelta(days=13)
        )

        chart_df["period"] = (
            chart_df["period_start"]
            .dt.strftime("%d %b")
            + " – "
            + period_end.dt.strftime("%d %b %Y")
        )

    return chart_df


# ---------------------------------------------------------------------------
# EXCEL EXPORT
# ---------------------------------------------------------------------------

def build_excel_report(summary_df):
    export_columns = [
        "Student",
        "Issues created",
        "Capture consistency",
        "Capture grade",
        "Average per period",
        "Periods with activity",
        "Email",
        "Section",
        "Jira URL",
    ]

    export_df = summary_df[export_columns].copy()

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        export_df.to_excel(
            writer,
            index=False,
            sheet_name="Capture Report",
        )

        worksheet = writer.sheets["Capture Report"]

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="E9ECEF",
        )

        thin_gray = Side(
            style="thin",
            color="BFC5CC",
        )

        header_border = Border(
            left=thin_gray,
            right=thin_gray,
            top=thin_gray,
            bottom=thin_gray,
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = header_border

        widths = {
            "A": 28,
            "B": 15,
            "C": 20,
            "D": 15,
            "E": 18,
            "F": 22,
            "G": 28,
            "H": 34,
            "I": 35,
        }

        for col, width in widths.items():
            worksheet.column_dimensions[col].width = width

        # Capture grade is numeric.
        grade_col = (
            export_df.columns.get_loc("Capture grade")
            + 1
        )

        for row_num in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_num,
                column=grade_col,
            ).number_format = "0.00"

        # Keep the exact dashboard fraction, e.g. "8 / 17".
        periods_col = (
            export_df.columns.get_loc(
                "Periods with activity"
            )
            + 1
        )

        for row_num, value in enumerate(
            export_df["Periods with activity"],
            start=2,
        ):
            cell = worksheet.cell(
                row=row_num,
                column=periods_col,
            )
            cell.value = str(value)
            cell.number_format = "@"

        worksheet.freeze_panes = "A2"

    return output.getvalue()


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.title("Jira Student Capture Dashboard")
st.caption(
    "Compare task-capture consistency across student Jira sites."
)

roster_path = find_roster_file()

if roster_path is None:
    st.error(
        "Student roster not found. Put student_repo.xlsx in the same "
        "folder as this dashboard script."
    )
    st.stop()

try:
    roster_df = load_student_roster(
        str(roster_path)
    )
except Exception as exc:
    st.error(f"Could not load student roster: {exc}")
    st.stop()


with st.sidebar:
    st.header("Jira credentials")

    default_email = os.getenv(
        "JIRA_EMAIL",
        "",
    )
    default_token = os.getenv(
        "JIRA_API_TOKEN",
        "",
    )

    jira_email = st.text_input(
        "Your Atlassian email",
        value=default_email,
        placeholder="you@example.com",
    )

    jira_token = st.text_input(
        "API token",
        value=default_token,
        type="password",
    )

    st.caption(
        f"Roster loaded automatically: "
        f"{len(roster_df)} students"
    )


st.subheader("Select students")

student_labels = {
    f"{row['Name']} — {row['Email']}": row["student_key"]
    for _, row in roster_df.iterrows()
}

selected_student_labels = st.multiselect(
    "Students",
    options=list(student_labels.keys()),
    placeholder=(
        "Choose one student, or select several to compare"
    ),
    help=(
        "Each student's Jira URL is read automatically "
        "from the Excel roster."
    ),
)

selected_student_keys = [
    student_labels[label]
    for label in selected_student_labels
]

selected_students_df = roster_df[
    roster_df["student_key"].isin(
        selected_student_keys
    )
].copy()


filter_col1, filter_col2, filter_col3 = st.columns(3)

today = date.today()
default_start = today - timedelta(days=29)

with filter_col1:
    start_date = st.date_input(
        "Start date",
        value=default_start,
        max_value=today,
    )

with filter_col2:
    end_date = st.date_input(
        "End date",
        value=today,
        max_value=today,
    )

with filter_col3:
    granularity = st.selectbox(
        "Time grouping",
        [
            "Daily",
            "Weekly",
            "Biweekly",
        ],
        index=1,
    )

if start_date > end_date:
    st.error(
        "Start date must be before or equal to end date."
    )
    st.stop()


issue_type_filter = st.selectbox(
    "Issue type",
    [
        "All issue types",
        "Task",
        "Story",
        "Bug",
        "Epic",
        "Sub-task",
    ],
    index=0,
    help=(
        "Use All issue types if you want all captured work. "
        "Choose Task if grading should count only Jira Tasks."
    ),
)

generate_report = st.button(
    "Generate report",
    type="primary",
)


if generate_report:
    if not jira_email or not jira_token:
        st.error(
            "Enter your Atlassian email and API token."
        )
        st.stop()

    if selected_students_df.empty:
        st.error(
            "Select at least one student."
        )
        st.stop()

    jql = build_jql(
        start_date,
        end_date,
        issue_type_filter,
    )

    frames = []
    failures = []

    progress = st.progress(0)
    status_text = st.empty()

    total_students = len(selected_students_df)

    for position, (_, student) in enumerate(
        selected_students_df.iterrows(),
        start=1,
    ):
        student_name = student["Name"]

        status_text.write(
            f"Reading {student_name} "
            f"({position}/{total_students})…"
        )

        try:
            student_df = fetch_student_issues(
                email=jira_email,
                api_token=jira_token,
                student_name=student_name,
                student_email=student["Email"],
                section=student["Section"],
                jira_url=student["Jira URL"],
                jql=jql,
            )

            if not student_df.empty:
                frames.append(student_df)

        except Exception as exc:
            failures.append(
                {
                    "Student": student_name,
                    "Jira URL": student["Jira URL"],
                    "Error": str(exc),
                }
            )

        progress.progress(
            position / total_students
        )

    status_text.empty()
    progress.empty()

    if frames:
        issues_df = pd.concat(
            frames,
            ignore_index=True,
        )
    else:
        issues_df = pd.DataFrame()

    st.session_state[
        "report_issues_df"
    ] = issues_df

    st.session_state[
        "report_students_df"
    ] = selected_students_df

    st.session_state[
        "report_failures"
    ] = failures

    st.session_state[
        "report_start_date"
    ] = start_date

    st.session_state[
        "report_end_date"
    ] = end_date

    st.session_state[
        "report_granularity"
    ] = granularity


if "report_students_df" not in st.session_state:
    st.info(
        "Choose one or more students and click "
        "**Generate report**."
    )
    st.stop()


issues_df = st.session_state["report_issues_df"]
report_students_df = st.session_state[
    "report_students_df"
]
report_failures = st.session_state.get(
    "report_failures",
    [],
)
report_start_date = st.session_state[
    "report_start_date"
]
report_end_date = st.session_state[
    "report_end_date"
]
report_granularity = st.session_state[
    "report_granularity"
]


if report_failures:
    with st.expander(
        f"{len(report_failures)} site(s) could not be read"
    ):
        st.dataframe(
            pd.DataFrame(report_failures),
            use_container_width=True,
            hide_index=True,
        )


successful_names = set(
    report_students_df["Name"]
)

if report_failures:
    failed_names = {
        item["Student"]
        for item in report_failures
    }

    successful_students_df = (
        report_students_df[
            ~report_students_df["Name"].isin(
                failed_names
            )
        ].copy()
    )
else:
    successful_students_df = (
        report_students_df.copy()
    )


summary_df = build_student_summary(
    issues_df,
    successful_students_df,
    report_start_date,
    report_end_date,
    report_granularity,
)

chart_df = build_chart_data(
    issues_df,
    successful_students_df,
    report_start_date,
    report_end_date,
    report_granularity,
)


# ---------------------------------------------------------------------------
# HEADLINE METRICS
# ---------------------------------------------------------------------------

if len(summary_df) == 1:
    row = summary_df.iloc[0]

    m1, m2, m3, m4 = st.columns(4)

    m1.metric(
        "Issues created",
        f"{int(row['Issues created']):,}",
    )

    m2.metric(
        "Capture consistency",
        row["Capture consistency"],
    )

    m3.metric(
        "Average per period",
        f"{float(row['Average per period']):.1f}",
    )

    m4.metric(
        "Periods with activity",
        row["Periods with activity"],
    )

elif len(summary_df) > 1:
    mean_consistency = (
        summary_df["_consistency_numeric"].mean()
    )
    mean_grade = (
        summary_df["Capture grade"].mean()
    )
    total_issues = int(
        summary_df["Issues created"].sum()
    )

    m1, m2, m3, m4 = st.columns(4)

    m1.metric(
        "Students compared",
        f"{len(summary_df)}",
    )

    m2.metric(
        "Issues created",
        f"{total_issues:,}",
    )

    m3.metric(
        "Average consistency",
        f"{mean_consistency:.1f}%",
    )

    m4.metric(
        "Average capture grade",
        f"{mean_grade:.2f} / 4",
    )


# ---------------------------------------------------------------------------
# CHART
# ---------------------------------------------------------------------------

if not chart_df.empty:
    if len(successful_students_df) > 1:
        fig = px.bar(
            chart_df,
            x="period",
            y="issues_created",
            color="student_name",
            barmode="group",
            labels={
                "period": report_granularity,
                "issues_created": "Issues created",
                "student_name": "Student",
            },
            title=(
                "Issues created — "
                f"{report_granularity.lower()} "
                "comparison"
            ),
        )
    else:
        fig = px.bar(
            chart_df,
            x="period",
            y="issues_created",
            labels={
                "period": report_granularity,
                "issues_created": "Issues created",
            },
            title=(
                "Issues created — "
                f"{report_granularity.lower()}"
            ),
        )

    fig.update_layout(
        xaxis_title=None,
        yaxis_title="Issues created",
        hovermode="x unified",
    )

    st.plotly_chart(
        fig,
        use_container_width=True,
    )


# ---------------------------------------------------------------------------
# STUDENT COMPARISON TABLE
# ---------------------------------------------------------------------------

st.markdown("### Student capture report")

display_columns = [
    "Student",
    "Issues created",
    "Capture consistency",
    "Capture grade",
    "Average per period",
    "Periods with activity",
]

st.dataframe(
    summary_df[display_columns],
    use_container_width=True,
    hide_index=True,
)


# ---------------------------------------------------------------------------
# DOWNLOAD
# ---------------------------------------------------------------------------

if not summary_df.empty:
    export_columns = [
        "Student",
        "Issues created",
        "Capture consistency",
        "Capture grade",
        "Average per period",
        "Periods with activity",
        "Email",
        "Section",
        "Jira URL",
    ]

    csv_report = (
        summary_df[export_columns]
        .to_csv(index=False)
        .encode("utf-8-sig")
    )

    excel_report = build_excel_report(
        summary_df
    )

    download_col1, download_col2 = st.columns(2)

    with download_col1:
        st.download_button(
            "Download CSV report",
            data=csv_report,
            file_name="jira_student_capture_report.csv",
            mime="text/csv",
            use_container_width=True,
        )

    with download_col2:
        st.download_button(
            "Download Excel report",
            data=excel_report,
            file_name="jira_student_capture_report.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )
